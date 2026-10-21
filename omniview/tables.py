"""XLSX / CSV table model helpers."""
from __future__ import annotations

import csv
import datetime as _dt
import io
import os
import re
from decimal import Decimal

from . import textutil


def _safe(v):
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v.isoformat(sep=" ") if isinstance(v, _dt.datetime) else v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, bytes):
        return v[:32].hex().upper()
    return str(v)


def _color(c):
    try:
        if c is None:
            return None
        if c.type == "rgb" and isinstance(c.rgb, str):
            rgb = c.rgb
            return "#" + (rgb[2:] if len(rgb) == 8 else rgb)
        if c.type == "theme":
            return "#5b7cfa"
    except Exception:
        pass
    return None


def _style(cell) -> dict:
    s = {}
    try:
        f = cell.font
        s.update(b=bool(f.bold), i=bool(f.italic), u=bool(f.underline), fn=f.name,
                 fs=float(f.size) if f.size else None, fc=_color(f.color))
    except Exception:
        pass
    try:
        fill = cell.fill
        if fill is not None and fill.fill_type not in (None, "none"):
            s["bg"] = _color(fill.fgColor)
    except Exception:
        pass
    try:
        al = cell.alignment
        if al:
            s.update(ha=al.horizontal, va=al.vertical, wrap=bool(al.wrap_text))
    except Exception:
        pass
    try:
        if cell.number_format and cell.number_format != "General":
            s["fmt"] = cell.number_format
    except Exception:
        pass
    return s


def _fmt(v, fmt):
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v.strftime("%Y-%m-%d %H:%M" if isinstance(v, _dt.datetime) and (v.hour or v.minute) else "%Y-%m-%d")
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)) and fmt:
        try:
            if "%" in fmt:
                dec = len(fmt.split(".")[1].split("%")[0]) if "." in fmt else 0
                return f"{v*100:.{dec}f}%"
            if any(x in fmt for x in ("¥", "$", "€", "£")):
                sym = next(x for x in ("¥", "$", "€", "£") if x in fmt)
                return f"{sym}{v:,.2f}" if "." in fmt else f"{sym}{v:,.0f}"
            if fmt.startswith("#,##0"):
                return f"{v:,.2f}" if ".00" in fmt else f"{v:,.0f}"
            if fmt in ("0.00", "0.0", "0.000"):
                return f"{v:.{len(fmt.split('.')[1])}f}"
        except Exception:
            pass
    if isinstance(v, float):
        return str(int(round(v))) if abs(v - round(v)) < 1e-9 else f"{v:.10g}"
    return str(v)


def xlsx_sheets(path: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False, read_only=False)
    wbv = None
    try:
        wbv = load_workbook(path, data_only=True, read_only=False)
    except Exception:
        pass
    sheets = []
    for ws in wb.worksheets:
        merges = [str(r) for r in ws.merged_cells.ranges]
        widths = {}
        for k, d in list(ws.column_dimensions.items())[:200]:
            if d.width:
                widths[k] = round(float(d.width), 2)
        sheets.append({"name": ws.title, "rows": ws.max_row or 0, "cols": ws.max_column or 0,
                       "merges": merges, "widths": widths, "freeze": str(ws.freeze_panes or "")})
    return {"sheets": sheets, "_wb": wb, "_wbv": wbv}


_CELL = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_FUNCS = {
    "SUM": lambda xs: sum(xs), "AVERAGE": lambda xs: sum(xs) / len(xs) if xs else 0,
    "MIN": lambda xs: min(xs) if xs else 0, "MAX": lambda xs: max(xs) if xs else 0,
    "COUNT": lambda xs: len(xs), "ABS": lambda xs: abs(xs[0]) if xs else 0,
    "ROUND": lambda xs: round(xs[0], int(xs[1]) if len(xs) > 1 else 0),
}


def _eval_formula(ws, formula: str, cache: dict, depth: int = 0):
    if depth > 6:
        return None
    expr = formula[1:].strip().replace("^", "**")

    def value(ref):
        ref = ref.replace("$", "").upper()
        if ref in cache:
            return cache[ref]
        cell = ws[ref]
        v = cell.value
        if v is None:
            cache[ref] = None
            return None
        if isinstance(v, str) and v.startswith("="):
            cache[ref] = None
            v = _eval_formula(ws, v, cache, depth + 1)
            cache[ref] = v
            return v
        if isinstance(v, bool):
            return 1.0 if v else 0.0
        if isinstance(v, (int, float, Decimal)):
            return float(v)
        return None

    def range_values(ref):
        m = re.match(r"^\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)$", ref.strip())
        if not m:
            v = value(ref)
            return [v] if v is not None else []
        c1, r1, c2, r2 = (_col_num(m.group(1)), int(m.group(2)), _col_num(m.group(3)), int(m.group(4)))
        out = []
        for r in range(min(r1, r2), min(max(r1, r2), 200) + 1):
            for c in range(min(c1, c2), min(max(c1, c2), 80) + 1):
                v = value(_col_name(c) + str(r))
                if v is not None:
                    out.append(v)
        return out

    for _ in range(8):
        m = re.search(r"\b(SUM|AVERAGE|MIN|MAX|COUNT|ABS|ROUND)\s*\(([^()]*)\)", expr, re.I)
        if not m:
            break
        fn = m.group(1).upper()
        vals = []
        args = [a.strip() for a in m.group(2).split(",") if a.strip()]
        for a in args:
            if re.match(r"^\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?$", a, re.I):
                vals.extend(range_values(a))
            else:
                try:
                    vals.append(float(a.rstrip("%")) / (100 if a.endswith("%") else 1))
                except ValueError:
                    return None
        try:
            res = _FUNCS[fn](vals)
        except Exception:
            return None
        expr = expr[:m.start()] + repr(res) + expr[m.end():]

    def repl(m):
        v = value(m.group(0))
        if v is None:
            raise ValueError
        return repr(float(v))

    try:
        expr = _CELL.sub(repl, expr)
    except ValueError:
        return None
    expr = expr.replace(",", "")
    if not re.match(r"^[0-9eE+\-*/().%\s]+$", expr):
        return None
    try:
        return float(eval(expr, {"__builtins__": None}, {}))
    except Exception:
        return None


def xlsx_data(path: str, sheet: str | None = None) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False, read_only=False)
    try:
        wbv = load_workbook(path, data_only=True, read_only=False)
    except Exception:
        wbv = None
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else wb.active
        wsv = wbv[ws.title] if wbv is not None and ws.title in wbv.sheetnames else None
        rows = min(ws.max_row or 1, 500)
        cols = min(ws.max_column or 1, 120)
        grid = []
        styles = {}
        style_list = []
        cache = {}
        for r in range(1, rows + 1):
            line = []
            for c in range(1, cols + 1):
                cell = ws.cell(row=r, column=c)
                v = cell.value
                f = None
                if isinstance(v, str) and v.startswith("="):
                    f = v
                    v = None
                    if wsv is not None:
                        cv = wsv.cell(row=r, column=c).value
                        if cv is not None:
                            v = cv
                    if v is None:
                        v = _eval_formula(ws, f, cache)
                fmt = cell.number_format if cell.number_format != "General" else None
                s = _style(cell)
                key = repr(sorted(s.items()))
                if key not in styles:
                    styles[key] = len(style_list)
                    style_list.append(s)
                line.append({"v": _safe(v) if v is not None else "", "f": f,
                             "d": _fmt(v, fmt) if v is not None else (f or ""),
                             "s": styles[key], "fmt": fmt})
            grid.append(line)
        return {"sheet": ws.title, "sheets": wb.sheetnames, "grid": grid,
                "styles": style_list, "rows": rows, "cols": cols,
                "merges": [str(m) for m in ws.merged_cells.ranges],
                "widths": {k: round(float(d.width), 2) for k, d in list(ws.column_dimensions.items())[:200] if d.width},
                "freeze": str(ws.freeze_panes or "")}
    finally:
        wb.close()
        if wbv is not None:
            wbv.close()


def csv_data(path: str) -> dict:
    raw = open(path, "rb").read()
    dec = textutil.decode_bytes(raw)
    text = dec["text"]
    delim = ","
    try:
        delim = csv.Sniffer().sniff(text[:32768], delimiters=",;\t|").delimiter
    except Exception:
        counts = {d: text.count(d) for d in [",", ";", "\t", "|"]}
        delim = max(counts, key=counts.get) if max(counts.values()) else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    cols = max((len(r) for r in rows), default=0)
    grid = []
    for r in rows[:2000]:
        grid.append([{"v": r[c] if c < len(r) else "", "f": None,
                      "d": r[c] if c < len(r) else "", "s": 0, "fmt": None} for c in range(cols)])
    return {"sheet": os.path.basename(path), "sheets": [os.path.basename(path)], "grid": grid,
            "styles": [{}], "rows": len(grid), "cols": cols, "merges": [], "widths": {},
            "freeze": "", "encoding": dec["label"], "delimiter": delim}
